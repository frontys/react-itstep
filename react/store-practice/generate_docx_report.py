import os
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_PATH = r"C:\AI\store-practice\ОТЧЕТ_Учёт_товаров_магазина.docx"
SCREENSHOTS_DIR = r"C:\AI\store-practice\screenshots"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report():
    doc = Document()

    # Поля страницы
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Заголовок
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("ОТЧЁТ О ПРАКТИЧЕСКОЙ РАБОТЕ\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    run_sub = p_title.add_run("«Учёт товаров магазина» (React + Express + SQLite)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph()

    # Раздел 1: Описание проекта
    h1 = doc.add_heading("1. Стек технологий и архитектура", level=1)
    h1.style.font.name = "Arial"

    p_tech = doc.add_paragraph()
    p_tech.add_run("• Клиент: ").bold = True
    p_tech.add_run("React 19, Vite, функциональные компоненты, хуки useState / useEffect, Fetch API, адаптивный CSS3.\n")
    p_tech.add_run("• Сервер: ").bold = True
    p_tech.add_run("Node.js 24.x, Express 5, CORS (порт 3001, origin http://localhost:5173).\n")
    p_tech.add_run("• База данных: ").bold = True
    p_tech.add_run("Встроенный в Node.js 24 модуль node:sqlite (класс DatabaseSync), файл server/store.db.\n")
    p_tech.add_run("• Архитектурный подход: ").bold = True
    p_tech.add_run("Разделение ответственности (SoC). Сервер обеспечивает строгую валидацию входных типов, безопасность SQL-запросов через подготовленные параметры (?) и изоляцию файловой системы базы данных от клиента.")

    # Раздел 2: Таблица 14 сценариев
    doc.add_heading("2. Результаты выполнения сценариев проверки", level=1)
    
    table_data = [
        ("№", "Действие", "Ожидаемый результат", "Фактический результат", "Статус"),
        ("1", "Первый запуск с новой базой", "Пустой список, доступна форма", "Сервер возвращает [], форма активна", "Пройден"),
        ("2", "Добавить 5 товаров со стр. 3 задания", "5 строк, ID различаются. Итого: 44 шт, 49 200 грн", "Добавлены 5 товаров. Сумма остатка: 44 шт, 49 200.00 грн", "Пройден"),
        ("3", "Обновить страницу браузера", "Все пять записей сохранены", "При повторном GET-запросе все 5 записей сохранены в SQLite", "Пройден"),
        ("4", "Изменить цену мыши на 700, остаток на 10", "Та же запись и ID; стоимость остатка 7000 грн", "Товар обновлён без изменения ID; 700 × 10 = 7000.00 грн", "Пройден"),
        ("5", "Начать изменение, затем нажать Отмена", "Форма очищена, база не изменена", "Поля формы очищены, editingId=null, база не затронута", "Пройден"),
        ("6", "Удалить клавиатуру и подтвердить", "Клавиатура удалена, осталось 4 товара", "DELETE успешно выполнен, в таблице и базе осталось 4 товара", "Пройден"),
        ("7", "Отказаться от удаления другого товара", "Товар остался в таблице", "confirm=false отменяет запрос, товар остался в базе", "Пройден"),
        ("8", "Остановить и запустить сервер", "Оставшиеся 4 товара доступны", "После полного перезапуска сервера все 4 товара прочитаны из store.db", "Пройден"),
        ("9", "Отправить POST с пустым name через API", "400 и JSON с message; запись не добавлена", "Получен код 400 и сообщение 'Введите название товара'", "Пройден"),
        ("10", "Отправить price = -1 или quantity = 1.5", "400; данные не сохранены", "Оба запроса отклонены сервером с кодом 400", "Пройден"),
        ("11", "Добавить товар с price = 0 и quantity = 0", "Добавление разрешено (код 201)", "Сервер вернул 201 Created; товар с 0/0 успешно создан", "Пройден"),
        ("12", "PUT и DELETE для несуществующего ID", "404 и понятное сообщение", "Сервер вернул статус 404: 'Товар не найден'", "Пройден"),
        ("13", "Передать id = abc", "400; сервер продолжает работать", "Сервер вернул статус 400: 'Некорректный ID', не упал", "Пройден"),
        ("14", "Остановить сервер и отправить форму", "Ошибка видна в UI, форма разблокирована", "Сетевая ошибка перехвачена, баннер ошибки показан, busy=false", "Пройден")
    ]

    table = doc.add_table(rows=len(table_data), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(0.4), Inches(1.8), Inches(2.0), Inches(2.0), Inches(0.8)]

    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            text = table_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.text = text
            p.runs[0].font.name = "Arial"
            p.runs[0].font.size = Pt(8.5)
            
            if row_idx == 0:
                p.runs[0].font.bold = True
                set_cell_background(cell, "F1F5F9")
            else:
                if col_idx == 4:
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = RGBColor(22, 101, 52)
                    set_cell_background(cell, "DCFCE7")
                elif row_idx % 2 == 1:
                    set_cell_background(cell, "FAFAFA")

    doc.add_paragraph()

    # Раздел 3: Скриншоты
    doc.add_heading("3. Скриншоты выполнения работы", level=1)

    screenshots = [
        ("Скриншот 1. Главный экран: сводка склада и таблица товаров", os.path.join(SCREENSHOTS_DIR, "screenshot1_list.png")),
        ("Скриншот 2. Заполнение формы добавления нового товара", os.path.join(SCREENSHOTS_DIR, "screenshot2_add.png")),
        ("Скриншот 3. Режим редактирования товара (кнопки сохранения и отмены)", os.path.join(SCREENSHOTS_DIR, "screenshot3_edit.png")),
        ("Скриншот 4. Отображение сообщения об ошибке валидации и сбоя сети", os.path.join(SCREENSHOTS_DIR, "screenshot4_error.png"))
    ]

    for title, img_path in screenshots:
        if os.path.exists(img_path):
            p_cap = doc.add_paragraph()
            p_cap.add_run(title).bold = True
            p_cap.runs[0].font.size = Pt(10)
            doc.add_picture(img_path, width=Inches(6.5))
            doc.add_paragraph()

    # Раздел 4: Контрольные вопросы
    doc.add_heading("4. Ответы на контрольные вопросы", level=1)

    questions_and_answers = [
        ("1. Зачем между React и SQLite нужен сервер?",
         "React выполняется на стороне клиента в изолированной среде браузера (sandbox). Браузер по соображениям безопасности не имеет доступа к локальной файловой системе и файлам СУБД SQLite. Прямой доступ к базе из клиента открывал бы критическую уязвимость: любой пользователь мог бы выполнить произвольный SQL-запрос. Сервер на Express выступает доверенным посредником: он проверяет права, производит строгую валидацию входных типов данных и безопасно выполняет SQL-запросы."),

        ("2. Чем состояние products отличается от таблицы products в базе?",
         "Состояние products в React хранится в оперативной памяти браузера в виде массива JavaScript-объектов. Оно существует только в пределах открытой вкладки и сбрасывается при обновлении страницы. Таблица products в SQLite хранится на постоянном дисковом накопителе сервера в файле store.db. Данные в базе персистентны и сохраняются при перезагрузке страницы, перезапуске сервера и выключении компьютера."),

        ("3. Зачем нужны JSON.stringify, express.json и response.json?",
         "• JSON.stringify() преобразует JavaScript-объект клиента в форматированную текстовую строку JSON для отправки по протоколу HTTP.\n"
         "• express.json() — это middleware сервера Express, который перехватывает входящие запросы с заголовком Content-Type: application/json, считывает поток байтов тела запроса, парсит JSON-строку и помещает готовый JavaScript-объект в request.body.\n"
         "• response.json() на сервере сериализует серверный объект/массив в формат JSON, устанавливает заголовок Content-Type: application/json и отправляет ответ клиенту."),

        ("4. Почему цену и количество нужно преобразовать из строк в числа?",
         "HTML-поля ввода возвращают значения через event.target.value всегда в виде строки (например, '650'). Если их не преобразовать в числа:\n"
         "1) В JavaScript операция сложения строк приведёт к конкатенации ('650' + '10' = '65010').\n"
         "2) Серверная валидация (typeof product.price !== 'number') отклонит запрос со статусом 400.\n"
         "3) В SQLite строковое значение может сохраниться в текстовом типе, что нарушит работу условий CHECK(price >= 0) и числовой сортировки."),

        ("5. Чем отличаются POST и PUT в этом приложении?",
         "• POST /api/products создаёт новую сущность: клиент не передаёт id, идентификатор генерируется базой (AUTOINCREMENT), сервер возвращает статус 201 Created с новым id.\n"
         "• PUT /api/products/:id обновляет все поля уже существующего товара с указанным id. Сервер проверяет result.changes и, если товар не найден, возвращает статус 404 Not Found."),

        ("6. Почему проверки React не заменяют проверки сервера?",
         "Проверки на клиенте в React предназначены исключительно для комфортного пользовательского опыта (UX) — мгновенно подсказать пользователю о незаполненном поле. Однако валидацию клиента элементарно обойти: злоумышленник или тестировщик может отправить HTTP-запрос напрямую через Postman, cURL или консоль браузера в обход формы. Сервер является единственным гарантированным рубежом защиты данных."),

        ("7. Как параметры ? помогают безопасно передавать значения в SQL?",
         "Параметры ? в подготовленных выражениях (Prepared Statements) отделяют синтаксис SQL-запроса от пользовательских данных. СУБД сначала компилирует структуру SQL-команды, а затем безопасно подставляет переданные значения как литералы. Это полностью исключает атаки типа SQL-инъекций (SQL Injection), так как кавычки и управляющие команды SQL не могут изменить логику запроса."),

        ("8. Как определить, что UPDATE или DELETE не нашёл запись?",
         "Свойство result.changes объекта, возвращаемого методом run() подготовленного запроса в node:sqlite, возвращает количество изменённых или удалённых строк. Если result.changes === 0, значит записи с таким id не существовало в базе, и сервер возвращает статус 404 Not Found. Если result.changes > 0, операция прошла успешно.")
    ]

    for q, a in questions_and_answers:
        p_q = doc.add_paragraph()
        p_q.add_run(q).bold = True
        p_q.runs[0].font.name = "Arial"
        p_q.runs[0].font.size = Pt(10.5)

        p_a = doc.add_paragraph()
        p_a.add_run(a)
        p_a.runs[0].font.name = "Arial"
        p_a.runs[0].font.size = Pt(9.5)

    # Раздел 5: Дополнительные задания
    doc.add_heading("5. Реализованные дополнительные задания", level=1)
    p_extra = doc.add_paragraph()
    p_extra.add_run("В проект успешно интегрированы следующие дополнительные улучшения:\n"
                    "1. Быстрый поиск в реальном времени по названию и категории товара.\n"
                    "2. Фильтр «Только нет в наличии» с моментальным отбором товаров с нулевым остатком.\n"
                    "3. Карточки суммарной аналитики склада: подсчёт общего количества товарных позиций, суммарного остатка единиц и совокупной денежной стоимости склада.")

    doc.save(DOC_PATH)
    print(f"Отчёт успешно сформирован: {DOC_PATH}")

if __name__ == "__main__":
    create_report()
